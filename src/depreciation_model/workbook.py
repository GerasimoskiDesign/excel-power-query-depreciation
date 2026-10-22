"""Build ``depreciation_model.xlsx`` with openpyxl.

openpyxl cannot embed Power Query, so the workbook ships with everything
*around* the queries: the parameters the M code reads (named ranges), the
source tables for inspection, the Python reference totals, an empty landing
sheet for the ``MonthlyTotals`` query and a ``Check`` sheet whose formulas
compare the two. Once the queries are pasted in (docs/LOAD_QUERIES.md) and
refreshed, the Check sheet turns green or shows exactly which month differs.
"""

from __future__ import annotations

import logging
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from depreciation_model.config import CHECK_METRICS, CHECK_TOLERANCE

log = logging.getLogger(__name__)

PLACEHOLDER_FOLDER = "C:\\CHANGE-ME\\excel-power-query-depreciation\\source_files\\"
PQ_SHEET = "PQ_MonthlyTotals"
EXPECTED_SHEET = "Expected_MonthlyTotals"
FIRST_DATA_ROW = 8
PQ_LAST_ROW = 2000

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")
AMBER = PatternFill("solid", fgColor="FFEB9C")
MONEY = "#,##0.00"


def _excel_value(v: object) -> object:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return None
    if isinstance(v, date) and not isinstance(v, datetime):
        return datetime(v.year, v.month, v.day)
    if hasattr(v, "item"):  # numpy scalar
        return v.item()
    return v


def _write_table(ws: Worksheet, df: pd.DataFrame, table_name: str) -> None:
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append([_excel_value(v) for v in row])
    for idx, col in enumerate(df.columns, start=1):
        letter = get_column_letter(idx)
        width = max(len(str(col)), 12)
        ws.column_dimensions[letter].width = min(width + 2, 40)
        if "date" in col or col == "month_start":
            for cell in ws[letter][1:]:
                cell.number_format = "yyyy-mm-dd"
        elif pd.api.types.is_float_dtype(df[col]):
            for cell in ws[letter][1:]:
                cell.number_format = MONEY
    if len(df):
        ref = f"A1:{get_column_letter(len(df.columns))}{len(df) + 1}"
        table = Table(displayName=table_name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
    ws.freeze_panes = "A2"


def _about_sheet(ws: Worksheet) -> None:
    lines = [
        ("Intangible-asset depreciation model", True),
        ("Synthetic data. Rebuilt from scratch; the production version is proprietary.", False),
        ("", False),
        ("How this workbook works", True),
        (
            "1. Parameters: set SourceFolder to the full path of the repo's source_files folder.",
            False,
        ),
        ("2. Paste the Power Query scripts from powerquery/ (see docs/LOAD_QUERIES.md).", False),
        ("3. Load query MonthlyTotals to sheet PQ_MonthlyTotals, cell A1.", False),
        ("4. Data > Refresh All. The Check sheet compares Power Query with the Python", False),
        ("   reference totals in Expected_MonthlyTotals (tolerance 0.01 per value).", False),
        ("", False),
        ("Sheets", True),
        ("AssetRegister - opening register (source data, as in source_files/).", False),
        ("Movements_Clean - movements as cleaned by the Python reference implementation.", False),
        ("Expected_MonthlyTotals - Python reference monthly totals (expected_outputs/).", False),
        ("PQ_MonthlyTotals - empty until the MonthlyTotals query is loaded here.", False),
        ("Check - formula comparison, one row per month.", False),
    ]
    for i, (text, bold) in enumerate(lines, start=1):
        ws.cell(row=i, column=1, value=text).font = Font(bold=bold, size=14 if i == 1 else 11)
    ws.column_dimensions["A"].width = 95


def _parameters_sheet(wb: Workbook, ws: Worksheet, source_folder: str, start: date, n: int) -> None:
    rows = [
        ("Parameter", "Value", "Used by"),
        (
            "SourceFolder",
            source_folder,
            "01_pSourceFolder.pq - folder with asset_register.csv and movements",
        ),
        (
            "ReportStart",
            datetime(start.year, start.month, 1),
            "02_pReportStart.pq - first reporting month",
        ),
        ("ReportMonths", n, "03_pReportMonths.pq - number of reporting months"),
    ]
    for r in rows:
        ws.append(list(r))
    for cell in ws[1]:
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
    ws["B3"].number_format = "yyyy-mm-dd"
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 55
    ws["A6"] = "Each value is a named range; Power Query reads it via Excel.CurrentWorkbook()."
    for name, ref in (("SourceFolder", "$B$2"), ("ReportStart", "$B$3"), ("ReportMonths", "$B$4")):
        wb.defined_names[name] = DefinedName(name, attr_text=f"Parameters!{ref}")


def _check_sheet(ws: Worksheet, expected: pd.DataFrame) -> None:
    n = len(expected)
    exp_cols = list(expected.columns)
    last = FIRST_DATA_ROW + n - 1

    ws["A1"] = "Check: Power Query MonthlyTotals vs Python reference (expected_outputs)"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"], ws["A3"], ws["A4"], ws["A5"] = (
        "Overall status",
        "PQ output loaded?",
        "Tolerance per value",
        "PQ columns in expected order?",
    )
    for c in ("A2", "A3", "A4", "A5"):
        ws[c].font = Font(bold=True)
    ws["B3"] = f"=LEN({PQ_SHEET}!$A$1)>0"
    ws["B4"] = CHECK_TOLERANCE
    header_checks = ",".join(
        f'{PQ_SHEET}!{get_column_letter(i)}1="{col}"' for i, col in enumerate(exp_cols, start=1)
    )
    ws["B5"] = f"=AND({header_checks})"

    ws.cell(row=7, column=1, value="month_start")
    status_col = 2 + 3 * len(CHECK_METRICS)
    for j, metric in enumerate(CHECK_METRICS):
        c0 = 2 + 3 * j
        ws.cell(row=6, column=c0, value=metric)
        for k, label in enumerate(("Expected", "Power Query", "Difference")):
            ws.cell(row=7, column=c0 + k, value=label)
    ws.cell(row=7, column=status_col, value="Row status")
    for row in (6, 7):
        for col in range(1, status_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.font, cell.fill = HEADER_FONT, HEADER_FILL
            cell.alignment = Alignment(horizontal="center")

    pq_month = f"{PQ_SHEET}!$A$2:$A${PQ_LAST_ROW}"
    for i in range(n):
        r = FIRST_DATA_ROW + i
        src = i + 2
        ws.cell(row=r, column=1, value=f"={EXPECTED_SHEET}!A{src}").number_format = "yyyy-mm"
        diffs = []
        for j, metric in enumerate(CHECK_METRICS):
            c0 = 2 + 3 * j
            letter = get_column_letter(exp_cols.index(metric) + 1)
            e, p, d = (get_column_letter(c0 + k) for k in range(3))
            ws[f"{e}{r}"] = f"={EXPECTED_SHEET}!{letter}{src}"
            pq_range = f"{PQ_SHEET}!${letter}$2:${letter}${PQ_LAST_ROW}"
            ws[f"{p}{r}"] = f'=IF($B$3,SUMIFS({pq_range},{pq_month},$A{r}),"")'
            ws[f"{d}{r}"] = f'=IF($B$3,{p}{r}-{e}{r},"")'
            for cell_ref in (e, p, d):
                ws[f"{cell_ref}{r}"].number_format = "#,##0" if metric == "asset_rows" else MONEY
            diffs.append(f"ABS({d}{r})<=$B$4")
        status = get_column_letter(status_col)
        ws[f"{status}{r}"] = (
            f'=IF(NOT($B$3),"NOT LOADED",IF(AND({",".join(diffs)}),"OK","MISMATCH"))'
        )

    status = get_column_letter(status_col)
    ws["B2"] = (
        '=IF(NOT($B$3),"NOT LOADED - load query MonthlyTotals to PQ_MonthlyTotals!A1",'
        f'IF(AND($B$5,COUNTIF({status}{FIRST_DATA_ROW}:{status}{last},"OK")={n}),'
        '"ALL CHECKS PASS","MISMATCH - see rows below"))'
    )
    ws["B2"].font = Font(bold=True)
    targets = ["B2:B2", f"{status}{FIRST_DATA_ROW}:{status}{last}"]
    for rng in targets:
        ws.conditional_formatting.add(
            rng, CellIsRule(operator="equal", formula=['"OK"'], fill=GREEN)
        )
        ws.conditional_formatting.add(
            rng, CellIsRule(operator="equal", formula=['"ALL CHECKS PASS"'], fill=GREEN)
        )
        ws.conditional_formatting.add(
            rng, CellIsRule(operator="equal", formula=['"MISMATCH"'], fill=RED)
        )
    ws.conditional_formatting.add(
        "B2:B2", CellIsRule(operator="notEqual", formula=['"ALL CHECKS PASS"'], fill=AMBER)
    )
    ws.column_dimensions["A"].width = 30
    for col in range(2, status_col + 1):
        ws.column_dimensions[get_column_letter(col)].width = 15
    ws.column_dimensions["B"].width = 34
    ws.freeze_panes = f"B{FIRST_DATA_ROW}"


def build_workbook(
    path: Path,
    register: pd.DataFrame,
    movements: pd.DataFrame,
    expected_totals: pd.DataFrame,
    report_start: date,
    report_months: int,
    source_folder: str = PLACEHOLDER_FOLDER,
) -> Path:
    """Write the workbook. ``expected_totals`` is :func:`reporting.monthly_totals` output."""
    wb = Workbook()
    _about_sheet(wb.active)
    wb.active.title = "About"
    _parameters_sheet(wb, wb.create_sheet("Parameters"), source_folder, report_start, report_months)
    _write_table(wb.create_sheet("AssetRegister"), register, "tblAssetRegister")
    _write_table(wb.create_sheet("Movements_Clean"), movements, "tblMovementsClean")
    expected = expected_totals.copy()
    for col in expected.columns:
        if pd.api.types.is_float_dtype(expected[col]):
            expected[col] = expected[col].round(2)
    _write_table(wb.create_sheet(EXPECTED_SHEET), expected, "tblExpectedMonthlyTotals")
    wb.create_sheet(PQ_SHEET)
    _check_sheet(wb.create_sheet("Check"), expected)
    wb.properties.creator = "depreciation_model"
    wb.properties.created = datetime(2026, 1, 1)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    log.info("Wrote workbook %s", path)
    return path
