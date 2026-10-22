"""The Excel Check sheet: evaluate its formulas with pycel for the three states it can be in."""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pycel import ExcelCompiler

from depreciation_model.config import CHECK_METRICS
from depreciation_model.pipeline import ModelResult
from depreciation_model.reporting import monthly_totals
from depreciation_model.workbook import EXPECTED_SHEET, FIRST_DATA_ROW, PQ_SHEET, build_workbook


@pytest.fixture(scope="module")
def workbook_path(tmp_path_factory: pytest.TempPathFactory, model: ModelResult) -> Path:
    path = tmp_path_factory.mktemp("wb") / "depreciation_model.xlsx"
    return build_workbook(
        path,
        model.register,
        model.movements,
        monthly_totals(model.schedule),
        date(2024, 1, 1),
        24,
        source_folder="C:\\data\\source_files\\",
    )


def _simulate_power_query_load(src: Path, dst: Path, tamper: tuple[int, str] | None = None) -> Path:
    """Copy the expected totals into PQ_MonthlyTotals, as a correct refresh would."""
    wb = load_workbook(src)
    expected, pq = wb[EXPECTED_SHEET], wb[PQ_SHEET]
    header = [c.value for c in expected[1]]
    for row in expected.iter_rows(values_only=True):
        pq.append(list(row))
    if tamper:
        row_idx, column = tamper
        cell = pq.cell(row=row_idx, column=header.index(column) + 1)
        cell.value = cell.value + 1.0
    wb.save(dst)
    return dst


def test_parameters_are_named_ranges(workbook_path: Path) -> None:
    wb = load_workbook(workbook_path)
    names = {
        n: wb.defined_names[n].attr_text for n in ("SourceFolder", "ReportStart", "ReportMonths")
    }
    assert names == {
        "SourceFolder": "Parameters!$B$2",
        "ReportStart": "Parameters!$B$3",
        "ReportMonths": "Parameters!$B$4",
    }
    assert wb["Parameters"]["B2"].value == "C:\\data\\source_files\\"
    assert wb["Parameters"]["B4"].value == 24


def test_workbook_contains_source_tables(workbook_path: Path, model: ModelResult) -> None:
    wb = load_workbook(workbook_path)
    assert wb["AssetRegister"].max_row == len(model.register) + 1
    assert wb["Movements_Clean"].max_row == len(model.movements) + 1
    assert "tblAssetRegister" in wb["AssetRegister"].tables


def test_check_sheet_says_not_loaded_before_power_query_runs(workbook_path: Path) -> None:
    xl = ExcelCompiler(filename=str(workbook_path))
    assert xl.evaluate("Check!B2").startswith("NOT LOADED")


def test_check_sheet_passes_when_power_query_matches(workbook_path: Path, tmp_path: Path) -> None:
    loaded = _simulate_power_query_load(workbook_path, tmp_path / "ok.xlsx")
    xl = ExcelCompiler(filename=str(loaded))
    assert xl.evaluate("Check!B5") is True
    assert xl.evaluate("Check!B2") == "ALL CHECKS PASS"


def test_check_sheet_flags_the_month_that_differs(workbook_path: Path, tmp_path: Path) -> None:
    # Row 5 of the PQ table = 4th month (2024-04). Off by 1.00 in depreciation.
    loaded = _simulate_power_query_load(
        workbook_path, tmp_path / "bad.xlsx", tamper=(5, "depreciation")
    )
    xl = ExcelCompiler(filename=str(loaded))
    status_col = chr(ord("A") + 1 + 3 * len(CHECK_METRICS))
    assert xl.evaluate("Check!B2").startswith("MISMATCH")
    assert xl.evaluate(f"Check!{status_col}{FIRST_DATA_ROW + 3}") == "MISMATCH"
    assert xl.evaluate(f"Check!{status_col}{FIRST_DATA_ROW + 2}") == "OK"


def test_check_sheet_flags_wrong_column_order(workbook_path: Path, tmp_path: Path) -> None:
    loaded = _simulate_power_query_load(workbook_path, tmp_path / "order.xlsx")
    wb = load_workbook(loaded)
    pq = wb[PQ_SHEET]
    pq["C1"], pq["D1"] = pq["D1"].value, pq["C1"].value  # swapped headers
    wb.save(loaded)
    xl = ExcelCompiler(filename=str(loaded))
    assert xl.evaluate("Check!B5") is False
    assert xl.evaluate("Check!B2").startswith("MISMATCH")
